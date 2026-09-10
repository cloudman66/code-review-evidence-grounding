from __future__ import annotations

import argparse
import csv
import json
import random
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx-path", default="benchmarks/cr_classification_assets/labeled_dataset.xlsx")
    parser.add_argument("--attributes-csv", default="benchmarks/cr_classification_assets/code_attributes.csv")
    parser.add_argument("--output-dir", default="data/processed/cr_intent")
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args()


def normalize_label(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def load_attributes(path: str) -> dict[str, dict[str, str]]:
    with Path(path).open("r", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return {row["folderName"]: row for row in reader}


def render_metric_context(attr_row: dict[str, str]) -> str:
    preferred = [
        "anyInserted",
        "anyDeleted",
        "getMovedSrcs",
        "UpdatedSrcs",
        "AddedOrUpdatedComments",
        "UpdatedFuncArguments",
        "numOldFiles",
        "numNewFiles",
    ]
    label_map = {
        "anyInserted": "inserted tokens",
        "anyDeleted": "deleted tokens",
        "getMovedSrcs": "moved sources",
        "UpdatedSrcs": "updated sources",
        "AddedOrUpdatedComments": "comments updated",
        "UpdatedFuncArguments": "updated function arguments",
        "numOldFiles": "old files",
        "numNewFiles": "new files",
    }
    fragments = []
    for key in preferred:
        value = attr_row.get(key, "")
        if value not in {"", "0", "-1", "-5"}:
            fragments.append(f"{label_map[key]}: {value}")
    return "; ".join(fragments) if fragments else "no notable structural code attributes"


def build_records(xlsx_path: str, attributes_csv: str) -> list[dict]:
    attr_map = load_attributes(attributes_csv)
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(min_row=2, values_only=True)

    records: list[dict] = []
    for row in rows:
        comment_id = row[0]
        message = row[8]
        comment_group = row[43]
        if not comment_id or not message or not comment_group:
            continue

        attr_row = attr_map.get(comment_id, {})
        records.append(
            {
                "sample_id": comment_id,
                "comment": str(message).strip(),
                "intent": normalize_label(comment_group),
                "gold_context_ids": [],
                "contexts": [
                    {
                        "context_id": "ctx-meta",
                        "source": "metadata",
                        "text": (
                            f"project: {row[1]}; file: {row[5]}; change_type: {row[6]}; "
                            f"line_number: {row[7]}; patchset_number: {row[19]}; "
                            f"is_bug_fix: {row[23]}; is_new_file: {row[24]}"
                        ),
                    },
                    {
                        "context_id": "ctx-metrics",
                        "source": "metrics",
                        "text": render_metric_context(attr_row),
                    },
                ],
                "metadata": {
                    "project": row[1],
                    "file_name": row[5],
                    "category_fine": row[14],
                    "comment_group": row[43],
                },
            }
        )

    return records


def stratified_split(records: list[dict], seed: int) -> dict[str, list[dict]]:
    rng = random.Random(seed)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        grouped[record["intent"]].append(record)

    splits = {"train": [], "dev": [], "test": []}
    for items in grouped.values():
        rng.shuffle(items)
        n = len(items)
        n_train = max(1, int(n * 0.7))
        n_dev = max(1, int(n * 0.1))
        if n_train + n_dev >= n:
            n_dev = 1
            n_train = max(1, n - 2)
        splits["train"].extend(items[:n_train])
        splits["dev"].extend(items[n_train:n_train + n_dev])
        splits["test"].extend(items[n_train + n_dev :])

    for split_records in splits.values():
        rng.shuffle(split_records)
    return splits


def write_jsonl(path: Path, records: list[dict]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    records = build_records(args.xlsx_path, args.attributes_csv)
    splits = stratified_split(records, seed=args.seed)

    for split_name, split_records in splits.items():
        write_jsonl(output_dir / f"{split_name}.jsonl", split_records)

    summary = {split_name: len(split_records) for split_name, split_records in splits.items()}
    (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
